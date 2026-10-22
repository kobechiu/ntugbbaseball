"""讀取 棒球訓練2026_整合.xlsx → 產生資料 JSON → 注入 template.html → 輸出 index.html。

用法: python3 tools/build.py [Excel 路徑]
"""
import json
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date
from pathlib import Path

import urllib.request

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "棒球訓練2026_整合.xlsx"
TEMPLATE = Path(__file__).resolve().parent / "template.html"
FORMS_CACHE = Path(__file__).resolve().parent / "forms_cache.json"
OUT = ROOT / "index.html"

# 網頁欄位 → Google 表單題目標題的關鍵字(依序比對)
FORM_FIELDS = {
    "回饋表單": {"name": ["名字", "姓名"], "item": ["項目"], "type": ["類型"], "content": ["想說", "建議", "內容"]},
    "新增表單": {"name": ["名稱"], "purpose": ["目的"], "part": ["部位"], "keyPoints": ["技術要點", "要點"],
                "video": ["影片"], "note": ["備註"]},
}

PURPOSE_FIX = {"外野手備": "外野守備"}
PART_FIX = {"髖": "下肢", "胸椎、手": "上肢", "軀幹": "全身"}
PARTS = ("上肢", "下肢", "全身")
issues = []


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def split_ids(v):
    return [x for x in re.split(r"[、,，\s]+", s(v)) if x]


def fix_part(v, label):
    p = PART_FIX.get(s(v), s(v))
    if p and p not in PARTS:
        issues.append(f"{label}:部位「{p}」不是 上肢/下肢/全身")
    return p


def fetch_form(url, keys):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    html = urllib.request.urlopen(req, timeout=20).read().decode("utf-8")
    m = re.search(r"FB_PUBLIC_LOAD_DATA_ = (.*?);</script>", html, re.S)
    if not m:
        raise ValueError("讀不到表單題目(表單是否設為需要登入?)")
    questions = []
    for q in json.loads(m.group(1))[1][1]:
        for sub in q[4] or []:
            opts = [o[0] for o in (sub[1] or []) if o and o[0]]
            questions.append({"title": q[1] or "", "entry": f"entry.{sub[0]}", "options": opts})
    fields, used = {}, set()
    for key, words in keys.items():
        for q in questions:
            if q["entry"] not in used and any(w in q["title"] for w in words):
                fields[key] = {"entry": q["entry"], "title": q["title"], "options": q["options"]}
                used.add(q["entry"])
                break
    for q in questions:
        if q["entry"] not in used:
            issues.append(f"表單題目「{q['title']}」網頁沒有對應欄位,送出時會留空")
    return {"action": re.sub(r"/viewform.*$", "/formResponse", url), "fields": fields}


def load_forms(wb):
    if "設定" not in wb.sheetnames:
        return {}
    conf = {s(r["項目"]): s(r["值"]) for r in rows(wb["設定"])}
    cache = json.loads(FORMS_CACHE.read_text(encoding="utf-8")) if FORMS_CACHE.exists() else {}
    forms = {}
    for name, keys in FORM_FIELDS.items():
        url = conf.get(name)
        if not url:
            continue
        try:
            forms[name] = fetch_form(url, keys)
        except Exception as e:
            if name not in cache:
                sys.exit(f"無法讀取{name}:{e}")
            issues.append(f"無法連線讀取{name}({e}),使用上次的表單設定")
            forms[name] = cache[name]
    FORMS_CACHE.write_text(json.dumps(forms, ensure_ascii=False, indent=1), encoding="utf-8")
    return forms


def rows(ws):
    header = [s(c.value) for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None and s(v) for v in r):
            yield dict(zip(header, r))


def parse_video(raw, label):
    v = s(raw)
    if not v:
        return []
    if not re.match(r"https?://", v):
        if re.match(r"(www\.|m\.)?[\w-]+\.(com|be)/", v):
            v = "https://" + v
        else:
            issues.append(f"{label}:影片欄位不是網址(「{v[:30]}…」)")
            return []
    if "/saved/" in v:
        issues.append(f"{label}:影片連結是個人「已收藏」頁面,其他人打不開")
    return [v]


def build():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    items, seen = [], defaultdict(int)

    for r in rows(wb["訓練項目"]):
        num = s(r.get("編號"))
        if not num:
            continue
        seen[num] += 1
        iid = num if seen[num] == 1 else f"{num}-{seen[num]}"
        name = s(r.get("名稱"))
        if seen[num] > 1:
            issues.append(f"編號 {num} 重複(「{name}」),網頁上暫時顯示為 {iid}")
        purpose = s(r.get("目的"))
        if purpose in PURPOSE_FIX:
            issues.append(f"#{num} 目的「{purpose}」視為「{PURPOSE_FIX[purpose]}」")
            purpose = PURPOSE_FIX[purpose]
        items.append({
            "id": iid, "num": num, "source": "教練", "name": name, "nameEn": "",
            "purpose": purpose, "part": fix_part(r.get("部位"), f"#{num}"), "keyPoints": s(r.get("技術要點")).lstrip(",，、 "),
            "videos": parse_video(r.get("影片連結"), f"#{num} {name}"), "note": s(r.get("備註")),
            "related": [],
        })

    for r in rows(wb["DL_動作"]):
        iid = s(r["編號"])
        items.append({
            "id": iid, "num": iid, "source": s(r["來源"]), "name": s(r["名稱"]), "nameEn": s(r["英文名稱"]),
            "purpose": s(r["目的"]), "part": fix_part(r["部位"], iid), "keyPoints": s(r["技術要點"]),
            "videos": parse_video(r["影片連結"], iid), "note": s(r["備註"]),
            "related": split_ids(r["相關項目"]),
        })

    by_num = defaultdict(list)
    for it in items:
        by_num[it["num"]].append(it["id"])

    def resolve(ids, where):
        out = []
        for n in ids:
            if n in by_num:
                out.extend(by_num[n])
            else:
                issues.append(f"{where}:找不到項目編號 {n}")
        return out

    menus = OrderedDict()
    for r in rows(wb["菜單"]):
        key = (s(r["來源表"]), s(r["對象"]), s(r["類別"]))
        m = menus.setdefault(key, {"sheet": key[0], "audience": key[1], "category": key[2], "steps": []})
        where = f"菜單 {key[0]}/{key[1]}/{s(r['步驟名稱'])}"
        entry = {"title": s(r["步驟名稱"]), "itemIds": resolve(split_ids(r["項目編號"]), where), "note": s(r["說明"])}
        label = s(r["分支"])
        if label:
            entry["label"] = label
            m["steps"][-1]["branches"].append(entry)
        else:
            entry["branches"] = []
            m["steps"].append(entry)

    equipment = OrderedDict()
    for r in rows(wb["器材"]):
        equipment[s(r["英文名稱"])] = {"zh": s(r["中文名稱"]), "desc": s(r["說明"]),
                                     "substitute": s(r["替代方案"]), "url": s(r["商品連結"])}
    env = {s(r["英文名稱"]): s(r["中文名稱"]) for r in rows(wb["環境"])}

    workouts = OrderedDict()
    for r in rows(wb["DL_課表類型"]):
        workouts[s(r["代碼"])] = {"code": s(r["代碼"]), "nameZh": s(r["中文名稱"]), "nameEn": s(r["英文名稱"]),
                                 "desc": s(r["說明"]), "color": s(r["顏色"]), "warmup": [], "skill": []}
    item_equipment = defaultdict(set)
    for r in rows(wb["DL_課表"]):
        w = workouts[s(r["課表代碼"])]
        did = s(r["動作編號"])
        if did not in by_num:
            issues.append(f"DL_課表 {w['code']}:找不到動作 {did}")
        row = {"itemId": did, "bat": s(r["球棒"]), "ball": s(r["球"]), "env": s(r["環境"]),
               "sets": s(r["組數"]), "reps": s(r["次數"]), "intensity": s(r["強度"])}
        for eq in (row["bat"], row["ball"]):
            if eq not in equipment:
                issues.append(f"DL_課表 {w['code']}:器材「{eq}」不在器材表")
            item_equipment[did].add(eq)
        w["warmup" if s(r["段落"]) == "熱身" else "skill"].append(row)
    for it in items:
        it["equipment"] = sorted(item_equipment.get(it["id"], []))

    cycles = OrderedDict()
    for r in rows(wb["DL_週期"]):
        c = cycles.setdefault(s(r["週期"]), {"name": s(r["週期"]), "weeks": []})
        days = [s(r[f"Day{i}"]) for i in range(1, 8)]
        for d in days:
            if d not in workouts and d not in ("休息", ""):
                issues.append(f"DL_週期 {c['name']} 第 {s(r['週'])} 週:未知課表代碼 {d}")
        c["weeks"].append({"week": s(r["週"]), "phase": s(r["階段"]), "days": days})

    articles = [{"id": s(r["編號"]), "title": s(r["標題"]), "titleEn": s(r["英文標題"]), "page": s(r["PDF頁"]),
                 "body": s(r["內容"]), "bodyEn": s(r["英文原文"])} for r in rows(wb["DL_說明"])]

    purposes = [s(r["目的"]) for r in rows(wb["詞條"]) if s(r.get("目的"))]

    return {
        "builtAt": date.today().isoformat(),
        "items": items, "menus": list(menus.values()), "workouts": list(workouts.values()),
        "cycles": list(cycles.values()), "articles": articles, "equipment": equipment, "env": env,
        "purposes": purposes, "forms": load_forms(wb),
    }


def main():
    data = build()
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8")
    if "/*__DATA__*/" not in html:
        sys.exit("template.html 缺少 /*__DATA__*/ 標記")
    OUT.write_text(html.replace("/*__DATA__*/", payload), encoding="utf-8")
    counts = defaultdict(int)
    for it in data["items"]:
        counts[it["source"]] += 1
    print(f"輸出 {OUT.name}:", dict(counts), f"菜單 {len(data['menus'])} 組,課表 {len(data['workouts'])} 種")
    if issues:
        print(f"\n資料檢查({len(issues)} 項):")
        for i in issues:
            print(" -", i)


if __name__ == "__main__":
    main()
